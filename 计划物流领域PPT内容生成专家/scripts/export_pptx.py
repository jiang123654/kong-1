"""
PPTX Export Script for 计划物流领域PPT内容生成专家
Version: v3.0
Author: PPT Content Generator

Usage:
    python export_pptx.py --input data.json --output output.pptx
"""

import json
import argparse
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
except ImportError:
    print("Error: python-pptx is required. Install with: pip install python-pptx")
    raise


class TCLColors:
    TCL_RED = RGBColor(0xE6, 0x00, 0x12)
    NAVY = RGBColor(0x1F, 0x29, 0x37)
    SUCCESS = RGBColor(0x10, 0xB9, 0x81)
    WARNING = RGBColor(0xF5, 0x9E, 0x0B)
    DANGER = RGBColor(0xEF, 0x44, 0x44)
    INFO = RGBColor(0x3B, 0x82, 0xF6)
    GRAY_50 = RGBColor(0xF9, 0xFA, 0xFB)
    GRAY_100 = RGBColor(0xF3, 0xF4, 0xF6)
    GRAY_200 = RGBColor(0xE5, 0xE7, 0xEB)
    GRAY_300 = RGBColor(0xD1, 0xD5, 0xDB)
    GRAY_400 = RGBColor(0x9C, 0xA3, 0xAF)
    GRAY_500 = RGBColor(0x6B, 0x72, 0x80)
    GRAY_600 = RGBColor(0x4B, 0x55, 0x63)
    GRAY_700 = RGBColor(0x37, 0x41, 0x51)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)


class PPTXExporter:
    def __init__(self, config_path=None):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.colors = TCLColors()
        self.config = self._load_config(config_path)
        self.slide_count = 0

    def _load_config(self, config_path):
        if config_path and Path(config_path).exists():
            with open(config_path, 'r', encoding='utf-8') as f:
                return json.load(f)
        return {}

    def _add_textbox(self, slide, left, top, width, height, text,
                     font_size=12, bold=False, color=None,
                     alignment=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                         Inches(width), Inches(height))
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor

        p = tf.paragraphs[0]
        p.alignment = alignment
        run = p.add_run()
        run.text = text
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.name = '微软雅黑'
        if color:
            run.font.color.rgb = color

        return txBox

    def _add_shape(self, slide, shape_type, left, top, width, height,
                   fill_color=None, line_color=None, line_width=None):
        shape = slide.shapes.add_shape(
            shape_type, Inches(left), Inches(top),
            Inches(width), Inches(height)
        )
        if fill_color:
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill_color
        else:
            shape.fill.background()

        if line_color:
            shape.line.color.rgb = line_color
            if line_width:
                shape.line.width = Pt(line_width)
        else:
            shape.line.fill.background()

        return shape

    def _add_rect(self, slide, left, top, width, height,
                  fill_color=None, line_color=None, line_width=None):
        return self._add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE,
                               left, top, width, height,
                               fill_color, line_color, line_width)

    def _add_footer(self, slide, page_num, total):
        self._add_textbox(slide, 0.5, 7.2, 4.0, 0.25,
                          "TCL实业 · 计划物流部",
                          font_size=9, color=self.colors.GRAY_400)
        self._add_textbox(slide, 12.0, 7.2, 0.83, 0.25,
                          f"{page_num} / {total}",
                          font_size=9, color=self.colors.GRAY_400,
                          alignment=PP_ALIGN.RIGHT)

    def add_cover_slide(self, title, subtitle="", presenter="", date=""):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.slide_count += 1

        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
            Inches(13.333), Inches(0.15)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = self.colors.TCL_RED
        shape.line.fill.background()

        self._add_textbox(slide, 1.0, 2.0, 11.3, 0.5,
                          "2025 H1总结 & H2规划",
                          font_size=16, bold=False,
                          color=self.colors.TCL_RED)

        self._add_textbox(slide, 1.0, 2.6, 11.3, 1.5,
                          title, font_size=40, bold=True,
                          color=self.colors.NAVY)

        self._add_textbox(slide, 1.0, 4.2, 11.3, 0.8,
                          subtitle, font_size=18,
                          color=self.colors.GRAY_500)

        meta_text = f"{presenter}   |   {date}"
        self._add_textbox(slide, 1.0, 5.5, 11.3, 0.5,
                          meta_text, font_size=14,
                          color=self.colors.GRAY_600)

        return slide

    def add_overview_slide(self, title, subtitle="", kpis=None, key_points=None):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.slide_count += 1
        kpis = kpis or []
        key_points = key_points or []

        self._add_page_header(slide, title, subtitle)

        card_width = 2.85
        card_height = 1.8
        card_top = 1.5
        gap = 0.2
        start_left = 0.8

        for i, kpi in enumerate(kpis[:4]):
            left = start_left + i * (card_width + gap)
            self._add_kpi_card(slide, left, card_top, card_width, card_height, kpi)

        if key_points:
            self._add_textbox(slide, 0.8, 3.8, 11.7, 0.4,
                              "核心要点", font_size=14, bold=True,
                              color=self.colors.NAVY)

            points_text = "\n".join([f"• {p}" for p in key_points])
            self._add_textbox(slide, 0.8, 4.3, 11.7, 2.5,
                              points_text, font_size=13,
                              color=self.colors.GRAY_700)

        self._add_footer(slide, self.slide_count, 0)
        return slide

    def _add_page_header(self, slide, title, subtitle=""):
        self._add_textbox(slide, 0.8, 0.4, 11.7, 0.5,
                          title, font_size=22, bold=True,
                          color=self.colors.NAVY)

        if subtitle:
            self._add_textbox(slide, 0.8, 0.9, 11.7, 0.4,
                              subtitle, font_size=12,
                              color=self.colors.GRAY_500)

        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.25),
            Inches(1.2), Inches(0.05)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = self.colors.TCL_RED
        line.line.fill.background()

    def _add_kpi_card(self, slide, left, top, width, height, kpi):
        status_color = {
            'success': self.colors.SUCCESS,
            'warning': self.colors.WARNING,
            'danger': self.colors.DANGER,
            'primary': self.colors.TCL_RED,
        }.get(kpi.get('status', 'primary'), self.colors.GRAY_300)

        card = self._add_rect(slide, left, top, width, height,
                              fill_color=self.colors.WHITE,
                              line_color=self.colors.GRAY_200, line_width=1)

        top_line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
            Inches(width), Inches(0.06)
        )
        top_line.fill.solid()
        top_line.fill.fore_color.rgb = status_color
        top_line.line.fill.background()

        self._add_textbox(slide, left + 0.15, top + 0.15, width - 0.3, 0.3,
                          kpi.get('label', ''), font_size=11,
                          color=self.colors.GRAY_500)

        self._add_textbox(slide, left + 0.15, top + 0.45, width - 0.3, 0.6,
                          kpi.get('value', ''), font_size=28, bold=True,
                          color=self.colors.NAVY)

        delta_text = kpi.get('delta', '')
        delta_color = self.colors.SUCCESS if kpi.get('deltaDirection', 'down') == 'down' else self.colors.DANGER
        compare_text = kpi.get('compare', '')

        self._add_textbox(slide, left + 0.15, top + 1.2, width - 0.3, 0.35,
                          f"{delta_text}  {compare_text}",
                          font_size=10, color=delta_color)

    def add_data_slide(self, title, subtitle="", chart_data=None,
                       table_data=None, insight=""):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.slide_count += 1

        self._add_page_header(slide, title, subtitle)

        if chart_data:
            self._add_bar_chart(slide, 0.8, 1.5, 7.0, 3.5, chart_data)

        if table_data:
            self._add_data_table(slide, 8.0, 1.5, 4.5, 3.5, table_data)

        if insight:
            self._add_insight_box(slide, 0.8, 5.2, 11.7, 0.9, insight)

        self._add_footer(slide, self.slide_count, 0)
        return slide

    def _add_bar_chart(self, slide, left, top, width, height, chart_data):
        chart_data_obj = CategoryChartData()
        chart_data_obj.categories = chart_data.get('categories', [])

        for series in chart_data.get('series', []):
            chart_data_obj.add_series(series.get('name', ''), series.get('data', []))

        chart_frame = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(left), Inches(top), Inches(width), Inches(height),
            chart_data_obj
        )
        chart = chart_frame.chart
        chart.has_title = False
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False

        plot = chart.plots[0]
        plot.gap_width = 100

        return chart

    def _add_data_table(self, slide, left, top, width, height, table_data):
        headers = table_data.get('headers', [])
        rows = table_data.get('rows', [])
        num_rows = len(rows) + 1
        num_cols = len(headers)

        table_shape = slide.shapes.add_table(
            num_rows, num_cols,
            Inches(left), Inches(top),
            Inches(width), Inches(height)
        )
        table = table_shape.table

        for j, header in enumerate(headers):
            cell = table.cell(0, j)
            cell.text = str(header)
            cell.fill.solid()
            cell.fill.fore_color.rgb = self.colors.GRAY_100
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(10)
                paragraph.font.bold = True
                paragraph.font.color.rgb = self.colors.GRAY_700
                paragraph.font.name = '微软雅黑'

        for i, row in enumerate(rows):
            for j, cell_val in enumerate(row):
                cell = table.cell(i + 1, j)
                if isinstance(cell_val, dict):
                    cell.text = str(cell_val.get('text', ''))
                else:
                    cell.text = str(cell_val)
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.size = Pt(10)
                    paragraph.font.color.rgb = self.colors.GRAY_700
                    paragraph.font.name = '微软雅黑'

        return table

    def _add_insight_box(self, slide, left, top, width, height, insight):
        box = self._add_rect(slide, left, top, width, height,
                             fill_color=RGBColor(0xFE, 0xF2, 0xF2),
                             line_color=self.colors.TCL_RED, line_width=1)

        self._add_textbox(slide, left + 0.2, top + 0.1, width - 0.4, height - 0.2,
                          f"关键洞察：{insight}",
                          font_size=11, color=self.colors.GRAY_700,
                          anchor=MSO_ANCHOR.MIDDLE)

    def add_rootcause_slide(self, title, subtitle="", layers=None):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.slide_count += 1
        layers = layers or []

        self._add_page_header(slide, title, subtitle)

        layer_width = 3.7
        layer_height = 4.5
        layer_top = 1.5
        gap = 0.25
        start_left = 0.9

        layer_colors = [
            (self.colors.INFO, RGBColor(0xEF, 0xF6, 0xFF)),
            (self.colors.WARNING, RGBColor(0xFF, 0xFB, 0xEB)),
            (self.colors.TCL_RED, RGBColor(0xFE, 0xF2, 0xF2)),
        ]

        for i, layer in enumerate(layers[:3]):
            left = start_left + i * (layer_width + gap)
            accent_color, bg_color = layer_colors[i]

            self._add_rootcause_layer(slide, left, layer_top, layer_width, layer_height,
                                      layer, i + 1, accent_color, bg_color)

        self._add_footer(slide, self.slide_count, 0)
        return slide

    def _add_rootcause_layer(self, slide, left, top, width, height,
                             layer, num, accent_color, bg_color):
        card = self._add_rect(slide, left, top, width, height,
                              fill_color=bg_color,
                              line_color=self.colors.GRAY_200, line_width=1)

        top_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
            Inches(width), Inches(0.08)
        )
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = accent_color
        top_bar.line.fill.background()

        header_text = f"0{num}  {layer.get('title', '')}"
        self._add_textbox(slide, left + 0.2, top + 0.2, width - 0.4, 0.4,
                          header_text, font_size=14, bold=True,
                          color=self.colors.NAVY)

        desc = layer.get('description', '')
        self._add_textbox(slide, left + 0.2, top + 0.7, width - 0.4, 0.5,
                          desc, font_size=11, color=self.colors.GRAY_700)

        items = layer.get('items', [])
        items_text = "\n".join([f"• {item}" for item in items])
        self._add_textbox(slide, left + 0.2, top + 1.4, width - 0.4, height - 2.0,
                          items_text, font_size=11, color=self.colors.GRAY_600)

        responsible = layer.get('responsible', '')
        if responsible:
            self._add_textbox(slide, left + 0.2, top + height - 0.5, width - 0.4, 0.35,
                              f"责任主体：{responsible}",
                              font_size=10, color=self.colors.GRAY_500)

    def add_cost_slide(self, title, subtitle="", chart_data=None,
                       table_data=None, summary=""):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.slide_count += 1

        self._add_page_header(slide, title, subtitle)

        if chart_data:
            self._add_bar_chart(slide, 0.8, 1.5, 7.0, 4.5, chart_data)

        if table_data:
            self._add_data_table(slide, 8.0, 1.5, 4.5, 4.5, table_data)

        if summary:
            self._add_insight_box(slide, 0.8, 6.2, 11.7, 0.7, summary)

        self._add_footer(slide, self.slide_count, 0)
        return slide

    def add_project_slide(self, title, subtitle="", milestones=None):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.slide_count += 1
        milestones = milestones or []

        self._add_page_header(slide, title, subtitle)

        if milestones:
            self._add_timeline(slide, 0.8, 1.8, 11.7, 4.0, milestones)

        self._add_footer(slide, self.slide_count, 0)
        return slide

    def _add_timeline(self, slide, left, top, width, height, milestones):
        num_items = len(milestones)
        if num_items == 0:
            return

        item_width = width / num_items
        line_y = top + 0.5

        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(left), Inches(line_y),
            Inches(width), Inches(0.04)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = self.colors.GRAY_300
        line.line.fill.background()

        for i, milestone in enumerate(milestones):
            x = left + i * item_width + item_width / 2

            dot = slide.shapes.add_shape(
                MSO_SHAPE.OVAL, Inches(x - 0.12), Inches(line_y - 0.1),
                Inches(0.24), Inches(0.24)
            )
            dot.fill.solid()
            dot.fill.fore_color.rgb = self.colors.TCL_RED
            dot.line.color.rgb = self.colors.WHITE
            dot.line.width = Pt(2)

            date_text = milestone.get('date', '')
            self._add_textbox(slide, x - item_width / 2, top, item_width, 0.4,
                              date_text, font_size=10, bold=True,
                              color=self.colors.TCL_RED,
                              alignment=PP_ALIGN.CENTER)

            title_text = milestone.get('title', '')
            self._add_textbox(slide, x - item_width / 2, line_y + 0.3,
                              item_width, 0.5,
                              title_text, font_size=11, bold=True,
                              color=self.colors.NAVY,
                              alignment=PP_ALIGN.CENTER)

            desc_text = milestone.get('description', '')
            self._add_textbox(slide, x - item_width / 2 + 0.1, line_y + 0.85,
                              item_width - 0.2, 1.5,
                              desc_text, font_size=10,
                              color=self.colors.GRAY_600,
                              alignment=PP_ALIGN.CENTER)

    def export(self, output_path):
        total = len(self.prs.slides)
        for i, slide in enumerate(self.prs.slides, 1):
            for shape in slide.shapes:
                if shape.has_text_frame:
                    if shape.text_frame.text == f"{{current}} / {{total}}":
                        shape.text_frame.text = f"{i} / {total}"

        self.prs.save(output_path)
        return output_path


def load_ppt_data(input_path):
    with open(input_path, 'r', encoding='utf-8') as f:
        return json.load(f)


def generate_pptx(data, output_path, config_path=None):
    exporter = PPTXExporter(config_path)

    meta = data.get('meta', {})
    slides_data = data.get('slides', [])

    for slide_data in slides_data:
        slide_type = slide_data.get('type', 'data')
        title = slide_data.get('title', '')
        subtitle = slide_data.get('subtitle', '')

        if slide_type == 'cover':
            exporter.add_cover_slide(
                title, subtitle,
                presenter=slide_data.get('presenter', ''),
                date=slide_data.get('date', '')
            )
        elif slide_type == 'overview':
            exporter.add_overview_slide(
                title, subtitle,
                kpis=slide_data.get('kpis', []),
                key_points=slide_data.get('keyPoints', [])
            )
        elif slide_type == 'rootcause':
            exporter.add_rootcause_slide(
                title, subtitle,
                layers=slide_data.get('layers', [])
            )
        elif slide_type == 'cost':
            exporter.add_cost_slide(
                title, subtitle,
                chart_data=slide_data.get('chartData'),
                table_data=slide_data.get('table'),
                summary=slide_data.get('summary', '')
            )
        elif slide_type == 'project':
            exporter.add_project_slide(
                title, subtitle,
                milestones=slide_data.get('milestones', [])
            )
        else:
            exporter.add_data_slide(
                title, subtitle,
                chart_data=slide_data.get('chartData'),
                table_data=slide_data.get('table'),
                insight=slide_data.get('insight', '')
            )

    exporter.export(output_path)
    return output_path


def main():
    parser = argparse.ArgumentParser(description='Export PPT content to PPTX')
    parser.add_argument('--input', required=True, help='Input JSON data file')
    parser.add_argument('--output', required=True, help='Output PPTX file path')
    parser.add_argument('--config', default=None, help='Layout config JSON path')
    args = parser.parse_args()

    data = load_ppt_data(args.input)
    output = generate_pptx(data, args.output, args.config)
    print(f"PPTX exported successfully: {output}")


if __name__ == '__main__':
    main()
